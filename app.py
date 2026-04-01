from __future__ import annotations

import hashlib
import json
import os
import re
from collections import Counter
from pathlib import Path
from urllib.parse import urlparse

import pandas as pd
import plotly.graph_objects as go
import requests
import streamlit as st

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import docx as python_docx
except ImportError:
    python_docx = None

try:
    from fpdf import FPDF
except ImportError:
    FPDF = None

DATA_PATH = Path("data/job_postings.csv")
SAMPLE_RESUME_PATH = Path("data/sample_resume.txt")
SAMPLE_RESUME_ALT_PATH = Path("data/sample_resume_alt.txt")
RESUME_VISSHVA_AIML_PATH = Path("data/resume_visshva_aiml_redacted.txt")
RESUME_VISSHVA_SDE_PATH = Path("data/resume_visshva_sde_redacted.txt")
LIVE_CACHE_PATH = Path("data/live_cache.csv")
SNAPSHOT_HISTORY_PATH = Path("data/profile_snapshots.json")

SKILL_CATALOG = [
    "Excel", "SQL", "Python", "Power BI", "Tableau", "DAX",
    "Data Cleaning", "Data Visualization", "Dashboard Storytelling",
    "Dashboard Design", "Data Modeling", "Statistics", "EDA",
    "Business Communication", "Communication", "Presentation Skills",
    "Business Analysis", "Documentation", "Reporting",
    "Stakeholder Reporting", "KPI Tracking", "Data Validation",
    "A/B Testing", "Experiment Analysis", "ETL Basics",
    "MIS Reporting", "Problem Solving", "Storytelling",
    "Attention to Detail", "LLM-assisted Analytics", "Prompt Engineering",
    "TypeScript", "Next.js", "Tailwind CSS", "Node.js", "Express.js",
    "REST APIs", "Docker", "Postman", "PostgreSQL", "MongoDB", "MySQL",
    "Prisma", "Supabase", "NextAuth.js", "HuggingFace Transformers",
    "DistilGPT-2", "Scikit-learn", "XGBoost", "NumPy", "Pandas", "Flask",
    "Java", "C++", "JavaScript", "HTML", "CSS", "React.js", "Git",
    "Machine Learning", "Deep Learning", "TensorFlow", "PyTorch",
    "NLP", "Computer Vision", "AWS", "Azure", "GCP", "Kubernetes",
    "CI/CD", "Agile", "Scrum", "JIRA", "Figma", "R",
]

SKILL_ALIAS = {
    "powerbi": "Power BI", "prompting": "Prompt Engineering",
    "llm analytics": "LLM-assisted Analytics", "story telling": "Storytelling",
    "nextjs": "Next.js", "tailwind": "Tailwind CSS",
    "typescript": "TypeScript", "node": "Node.js", "express": "Express.js",
    "postgres": "PostgreSQL", "postgresql": "PostgreSQL", "mongodb": "MongoDB",
    "scikit learn": "Scikit-learn", "sklearn": "Scikit-learn",
    "huggingface": "HuggingFace Transformers", "reactjs": "React.js",
    "react": "React.js", "ml": "Machine Learning", "dl": "Deep Learning",
    "nlp": "NLP", "cv": "Computer Vision", "amazon web services": "AWS",
    "google cloud": "GCP",
}

PROJECT_MAP = {
    "Power BI": "Build a Chennai retail KPI dashboard with filters, drill-down views, and one executive summary page.",
    "Dashboard Storytelling": "Create a business dashboard that ends with three written insights and one recommendation slide.",
    "DAX": "Build a Power BI report using calculated measures for revenue, growth, and customer retention.",
    "LLM-assisted Analytics": "Create an analytics copilot that turns natural language questions into chart suggestions from a CSV file.",
    "Prompt Engineering": "Design prompts for summarizing dataset insights and compare output quality using a fixed rubric.",
    "Data Modeling": "Create a star-schema style mini warehouse for sales data and show how joins improve reporting speed.",
    "A/B Testing": "Analyze a mock marketing campaign and recommend the winning variant using a simple significance check.",
    "TypeScript": "Convert an existing React UI to TypeScript with strict types and robust API error handling.",
    "Next.js": "Build a small Next.js app with SSR/SSG and an authenticated dashboard flow.",
    "Tailwind CSS": "Rebuild a UI with Tailwind and document a small design system (tokens + components).",
    "Docker": "Containerize a small API + UI, document local run steps, and add health checks.",
    "REST APIs": "Build and document a REST API with pagination, validation, and consistent error responses.",
    "HuggingFace Transformers": "Fine-tune or prompt a small transformer model and evaluate outputs with a measurable rubric.",
    "SQL": "Write advanced SQL queries with window functions, CTEs, and performance optimizations for a sample dataset.",
    "Python": "Build a CLI data pipeline that ingests, cleans, and summarizes CSV data with logging and error handling.",
    "Machine Learning": "Train and evaluate a classification model on a real-world dataset with proper cross-validation.",
    "React.js": "Build a responsive single-page app with state management, routing, and API integration.",
    "Node.js": "Create a RESTful backend with authentication, middleware, and database integration.",
    "Excel": "Build an advanced Excel workbook with pivot tables, lookup chains, and automated summary dashboards.",
    "Tableau": "Design an interactive Tableau dashboard with calculated fields, parameters, and storytelling.",
}

SKILL_CATEGORIES = {
    "Technical": ["Python", "SQL", "Java", "C++", "JavaScript", "TypeScript", "R", "HTML", "CSS"],
    "Data & Analytics": ["Excel", "Power BI", "Tableau", "Data Cleaning", "Data Visualization",
                         "Dashboard Storytelling", "Dashboard Design", "Data Modeling", "Statistics",
                         "EDA", "Data Validation", "A/B Testing", "Experiment Analysis"],
    "AI / ML": ["Machine Learning", "Deep Learning", "Scikit-learn", "XGBoost", "NumPy", "Pandas",
                "HuggingFace Transformers", "DistilGPT-2", "TensorFlow", "PyTorch", "NLP",
                "Computer Vision", "LLM-assisted Analytics", "Prompt Engineering"],
    "Web & Cloud": ["Next.js", "Tailwind CSS", "Node.js", "Express.js", "REST APIs", "Docker",
                    "React.js", "Flask", "Prisma", "Supabase", "NextAuth.js", "AWS", "Azure",
                    "GCP", "Kubernetes", "CI/CD"],
    "Soft Skills": ["Business Communication", "Communication", "Presentation Skills",
                    "Business Analysis", "Documentation", "Reporting", "Stakeholder Reporting",
                    "KPI Tracking", "Problem Solving", "Storytelling", "Attention to Detail",
                    "Agile", "Scrum"],
    "Tools": ["Git", "Postman", "JIRA", "Figma", "MongoDB", "MySQL", "PostgreSQL", "ETL Basics",
              "MIS Reporting", "DAX"],
}

LEARNING_RESOURCES = {
    "Power BI": [
        {"title": "Microsoft Power BI Learning Path", "url": "https://learn.microsoft.com/en-us/training/powerplatform/power-bi", "type": "Official Docs", "time": "10 hrs"},
        {"title": "Power BI Full Course - freeCodeCamp", "url": "https://www.youtube.com/watch?v=3u7MQz1EyPY", "type": "Video", "time": "4 hrs"},
    ],
    "SQL": [
        {"title": "SQLBolt - Interactive SQL Tutorials", "url": "https://sqlbolt.com/", "type": "Interactive", "time": "3 hrs"},
        {"title": "Mode Analytics SQL Tutorial", "url": "https://mode.com/sql-tutorial/", "type": "Tutorial", "time": "6 hrs"},
    ],
    "Python": [
        {"title": "Python Official Tutorial", "url": "https://docs.python.org/3/tutorial/", "type": "Official Docs", "time": "8 hrs"},
        {"title": "Automate the Boring Stuff", "url": "https://automatetheboringstuff.com/", "type": "Book (Free)", "time": "15 hrs"},
    ],
    "Dashboard Storytelling": [
        {"title": "Storytelling with Data", "url": "https://www.storytellingwithdata.com/", "type": "Book", "time": "5 hrs"},
        {"title": "Data Storytelling Guide", "url": "https://www.tableau.com/learn/articles/data-storytelling", "type": "Article", "time": "1 hr"},
    ],
    "Prompt Engineering": [
        {"title": "Prompt Engineering Guide", "url": "https://www.promptingguide.ai/", "type": "Guide", "time": "4 hrs"},
        {"title": "DeepLearning.AI Prompt Course", "url": "https://www.deeplearning.ai/short-courses/chatgpt-prompt-engineering-for-developers/", "type": "Course", "time": "2 hrs"},
    ],
    "Docker": [
        {"title": "Docker Getting Started", "url": "https://docs.docker.com/get-started/", "type": "Official Docs", "time": "3 hrs"},
    ],
    "TypeScript": [
        {"title": "TypeScript Handbook", "url": "https://www.typescriptlang.org/docs/handbook/", "type": "Official Docs", "time": "6 hrs"},
        {"title": "TypeScript Exercises", "url": "https://typescript-exercises.github.io/", "type": "Interactive", "time": "4 hrs"},
    ],
    "React.js": [
        {"title": "React Official Tutorial", "url": "https://react.dev/learn", "type": "Official Docs", "time": "6 hrs"},
    ],
    "Machine Learning": [
        {"title": "Google ML Crash Course", "url": "https://developers.google.com/machine-learning/crash-course", "type": "Course", "time": "15 hrs"},
    ],
    "Tableau": [
        {"title": "Tableau Training", "url": "https://www.tableau.com/learn/training", "type": "Official", "time": "8 hrs"},
    ],
    "Next.js": [
        {"title": "Next.js Learn Course", "url": "https://nextjs.org/learn", "type": "Official", "time": "6 hrs"},
    ],
}

MICRO_CURRICULUM_TEMPLATES = {
    "Power BI": [
        "Connect Power BI to a small relational dataset and clean fields for reporting.",
        "Create one KPI page, one drill-down page, and one stakeholder summary page.",
        "Write 3 business insights and 1 recommendation based on the dashboard.",
    ],
    "SQL": [
        "Practice SELECT, JOIN, GROUP BY, and filtering on a realistic analytics dataset.",
        "Write one CTE and one window-function query for trend analysis.",
        "Export your best 3 queries into a portfolio-ready SQL script with comments.",
    ],
    "Prompt Engineering": [
        "Design prompts for summarization, extraction, and reasoning over the same dataset.",
        "Compare outputs using a small evaluation rubric for accuracy and usefulness.",
        "Turn the best prompt set into a repeatable workflow with before/after examples.",
    ],
}

DEFAULT_GITHUB_SKILL_MAP = {
    "python": "Python",
    "jupyter notebook": "Python",
    "typescript": "TypeScript",
    "javascript": "JavaScript",
    "html": "HTML",
    "css": "CSS",
    "react": "React.js",
    "next.js": "Next.js",
    "nextjs": "Next.js",
    "node": "Node.js",
    "node.js": "Node.js",
    "express": "Express.js",
    "flask": "Flask",
    "sql": "SQL",
    "postgresql": "PostgreSQL",
    "mysql": "MySQL",
    "mongodb": "MongoDB",
    "docker": "Docker",
    "kubernetes": "Kubernetes",
    "aws": "AWS",
    "azure": "Azure",
    "gcp": "GCP",
    "pandas": "Pandas",
    "numpy": "NumPy",
    "scikit-learn": "Scikit-learn",
    "sklearn": "Scikit-learn",
    "tensorflow": "TensorFlow",
    "pytorch": "PyTorch",
    "nlp": "NLP",
    "computer-vision": "Computer Vision",
    "computer vision": "Computer Vision",
    "power bi": "Power BI",
    "tableau": "Tableau",
    "prompt-engineering": "Prompt Engineering",
    "prompt engineering": "Prompt Engineering",
    "rest-api": "REST APIs",
    "rest api": "REST APIs",
    "api": "REST APIs",
    "tailwindcss": "Tailwind CSS",
    "tailwind": "Tailwind CSS",
    "postman": "Postman",
    "git": "Git",
}

ROLE_QUERY_MAP = {
    "Junior Data Analyst": "junior data analyst",
    "Frontend Developer": "frontend developer",
    "AI/ML Intern": "machine learning intern",
    "Business Analyst": "business analyst",
    "Data Science Intern": "data science intern",
    "SDE / Full-stack Developer": "full stack developer",
}

CITY_TO_STATE_MAP = {
    "Chennai": "Tamil Nadu",
    "Bengaluru": "Karnataka",
    "Hyderabad": "Telangana",
    "Pune": "Maharashtra",
    "Gurugram": "Haryana",
}

CITY_ALIAS_MAP = {
    "bangalore": "Bengaluru",
    "bengaluru": "Bengaluru",
    "madras": "Chennai",
    "chennai": "Chennai",
    "hyderabad": "Hyderabad",
    "gurgaon": "Gurugram",
    "gurugram": "Gurugram",
}

CAREER_PATHS = {
    "Junior Data Analyst": {
        "current": "Junior Data Analyst",
        "paths": [
            {"role": "Senior Data Analyst", "years": "2-3 yrs", "key_skills": ["SQL", "Power BI", "Dashboard Storytelling", "Statistics"]},
            {"role": "Business Intelligence Analyst", "years": "2-4 yrs", "key_skills": ["Power BI", "DAX", "Data Modeling", "ETL Basics"]},
            {"role": "Data Scientist", "years": "3-5 yrs", "key_skills": ["Python", "Machine Learning", "Statistics", "NLP"]},
            {"role": "Analytics Manager", "years": "5-7 yrs", "key_skills": ["Business Analysis", "Stakeholder Reporting", "KPI Tracking", "Communication"]},
        ],
    },
    "Frontend Developer": {
        "current": "Frontend Developer",
        "paths": [
            {"role": "Senior Frontend Engineer", "years": "2-3 yrs", "key_skills": ["TypeScript", "Next.js", "React.js", "Tailwind CSS"]},
            {"role": "Full-Stack Developer", "years": "2-4 yrs", "key_skills": ["Node.js", "REST APIs", "PostgreSQL", "Docker"]},
            {"role": "Frontend Architect", "years": "4-6 yrs", "key_skills": ["TypeScript", "CI/CD", "Kubernetes", "Agile"]},
            {"role": "Engineering Manager", "years": "5-8 yrs", "key_skills": ["Agile", "Communication", "Scrum", "JIRA"]},
        ],
    },
    "SDE / Full-stack Developer": {
        "current": "SDE / Full-stack Developer",
        "paths": [
            {"role": "Senior SDE", "years": "2-3 yrs", "key_skills": ["Node.js", "TypeScript", "Docker", "REST APIs"]},
            {"role": "Backend Specialist", "years": "2-4 yrs", "key_skills": ["Node.js", "PostgreSQL", "Docker", "Kubernetes"]},
            {"role": "DevOps Engineer", "years": "3-5 yrs", "key_skills": ["Docker", "Kubernetes", "CI/CD", "AWS"]},
            {"role": "Tech Lead", "years": "4-6 yrs", "key_skills": ["Agile", "Communication", "Docker", "CI/CD"]},
        ],
    },
}

CUSTOM_CSS = """
<style>
    .main .block-container { padding-top: 1.5rem; max-width: 1200px; }
    div[data-testid="stMetric"] {
        background: linear-gradient(135deg, #667eea11, #764ba211);
        border: 1px solid #667eea33;
        border-radius: 12px;
        padding: 16px 20px;
        box-shadow: 0 2px 8px rgba(0,0,0,0.04);
    }
    div[data-testid="stMetric"] label {
        font-size: 0.85rem !important;
        font-weight: 600 !important;
        text-transform: uppercase;
        letter-spacing: 0.5px;
    }
    .stTabs [data-baseweb="tab-list"] { gap: 8px; }
    .stTabs [data-baseweb="tab"] {
        border-radius: 8px 8px 0 0;
        padding: 8px 20px;
        font-weight: 600;
    }
    .skill-tag {
        display: inline-block;
        padding: 4px 12px;
        margin: 3px;
        border-radius: 16px;
        font-size: 0.85rem;
        font-weight: 500;
    }
    .skill-matched { background-color: #d4edda; color: #155724; border: 1px solid #c3e6cb; }
    .skill-missing { background-color: #f8d7da; color: #721c24; border: 1px solid #f5c6cb; }
    .skill-rising { background-color: #fff3cd; color: #856404; border: 1px solid #ffeeba; }
    .section-header {
        font-size: 1.3rem;
        font-weight: 700;
        margin-bottom: 0.5rem;
        padding-bottom: 0.3rem;
        border-bottom: 2px solid #667eea44;
    }
    footer { visibility: hidden; }
</style>
"""


def normalize_token(text: str) -> str:
    return re.sub(r"[^a-z0-9+#]+", " ", text.lower()).strip()


def extract_text_from_pdf(uploaded_file) -> str:
    if pdfplumber is None:
        return "[PDF extraction unavailable — install pdfplumber]"
    text_parts = []
    with pdfplumber.open(uploaded_file) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_text_from_docx(uploaded_file) -> str:
    if python_docx is None:
        return "[DOCX extraction unavailable — install python-docx]"
    doc = python_docx.Document(uploaded_file)
    return "\n".join(para.text for para in doc.paragraphs if para.text.strip())


def extract_text_from_upload(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)
    if name.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)
    return uploaded_file.read().decode("utf-8", errors="replace")


def parse_github_username(raw_value: str) -> str:
    value = raw_value.strip()
    if not value:
        return ""
    if "github.com" not in value.lower():
        return value.lstrip("@").strip("/")
    parsed = urlparse(value)
    path_parts = [part for part in parsed.path.split("/") if part]
    return path_parts[0] if path_parts else ""


def get_secret(name: str, default: str = "") -> str:
    env_value = os.getenv(name)
    if env_value:
        return env_value
    try:
        return st.secrets.get(name, default)
    except Exception:
        return default


def _github_headers() -> dict[str, str]:
    token = get_secret("GITHUB_TOKEN")
    headers = {
        "Accept": "application/vnd.github+json",
        "User-Agent": "SkillPulse",
    }
    if token:
        headers["Authorization"] = f"Bearer {token}"
    return headers


@st.cache_data(show_spinner=False, ttl=1800)
def fetch_github_profile_data(username: str) -> dict[str, object]:
    if not username:
        return {"username": "", "skills": [], "summary": "", "repos": []}
    headers = _github_headers()
    base_url = f"https://api.github.com/users/{username}"
    user_response = requests.get(base_url, headers=headers, timeout=15)
    user_response.raise_for_status()
    repos_response = requests.get(f"{base_url}/repos?per_page=6&sort=updated", headers=headers, timeout=20)
    repos_response.raise_for_status()
    repos = repos_response.json()

    extracted_tokens: list[str] = []
    repo_summaries: list[str] = []
    for repo in repos:
        if repo.get("fork"):
            continue
        repo_name = repo.get("name", "")
        description = repo.get("description") or ""
        topics = repo.get("topics") or []
        language = repo.get("language") or ""
        extracted_tokens.extend([repo_name, description, language, *topics])
        repo_summaries.append(
            f"{repo_name}: {description or 'No description'}"
        )
    normalized_text = normalize_token(" ".join(extracted_tokens))
    detected: list[str] = []
    for token, canonical in DEFAULT_GITHUB_SKILL_MAP.items():
        if _skill_in_text(normalize_token(token), normalized_text) and canonical not in detected:
            detected.append(canonical)
    detected.extend([skill for skill in extract_skills(" ".join(extracted_tokens)) if skill not in detected])
    summary = user_response.json().get("bio") or ""
    return {
        "username": username,
        "skills": sorted(detected),
        "summary": summary,
        "repos": repo_summaries[:5],
    }


def _adzuna_secrets() -> tuple[str, str, str]:
    app_id = get_secret("ADZUNA_APP_ID")
    app_key = get_secret("ADZUNA_APP_KEY")
    country = get_secret("ADZUNA_COUNTRY", "in")
    return app_id, app_key, country


def _extract_api_skills(job_text: str) -> list[str]:
    skills = extract_skills(job_text)
    return skills[:8]


def _normalize_live_job_row(raw_job: dict, role: str, city: str) -> dict | None:
    title = raw_job.get("title") or role
    company_info = raw_job.get("company") or {}
    company = company_info.get("display_name") if isinstance(company_info, dict) else str(company_info or "Unknown Company")
    location_info = raw_job.get("location") or {}
    area = location_info.get("area") if isinstance(location_info, dict) else []
    detected_city = city
    if isinstance(area, list) and area:
        normalized_area = [str(part).strip() for part in area if str(part).strip()]
        city_matches = [part for part in normalized_area if part.lower() == city.lower()]
        if city_matches:
            detected_city = city_matches[0]
        else:
            state_name = CITY_TO_STATE_MAP.get(city, "").lower()
            non_state_parts = [part for part in normalized_area if part.lower() != state_name]
            detected_city = non_state_parts[-1] if non_state_parts else city
    description = raw_job.get("description") or ""
    skill_list = _extract_api_skills(f"{title}\n{description}")
    if not skill_list:
        return None
    posted_date = pd.to_datetime(raw_job.get("created"), errors="coerce")
    if pd.isna(posted_date):
        posted_date = pd.Timestamp.utcnow()
    return {
        "job_id": raw_job.get("id") or f"adzuna-{hashlib.md5((title + company).encode('utf-8')).hexdigest()[:10]}",
        "role": role,
        "city": detected_city or city,
        "posted_date": posted_date.strftime("%Y-%m-%d"),
        "title": title,
        "company": company,
        "skills": ";".join(skill_list),
    }


def fetch_adzuna_jobs(role: str, city: str, limit: int) -> pd.DataFrame:
    app_id, app_key, country = _adzuna_secrets()
    if not app_id or not app_key:
        raise RuntimeError("Adzuna credentials are missing. Add `ADZUNA_APP_ID` and `ADZUNA_APP_KEY` to enable live refresh.")
    query = ROLE_QUERY_MAP.get(role, role)
    state = CITY_TO_STATE_MAP.get(city, "")
    params = {
        "app_id": app_id,
        "app_key": app_key,
        "results_per_page": min(limit, 50),
        "what": query,
        "where": city,
        "content-type": "application/json",
    }
    if state:
        params["where"] = f"{city}, {state}"
    response = requests.get(
        f"https://api.adzuna.com/v1/api/jobs/{country}/search/1",
        params=params,
        timeout=25,
    )
    response.raise_for_status()
    results = response.json().get("results", [])
    rows = []
    for raw_job in results:
        normalized = _normalize_live_job_row(raw_job, role=role, city=city)
        if normalized:
            rows.append(normalized)
    if not rows:
        raise RuntimeError("The live API returned jobs, but none contained recognizable skills for the current role.")
    return pd.DataFrame(rows)


def load_jobs() -> pd.DataFrame:
    df = pd.read_csv(DATA_PATH)
    df["posted_date"] = pd.to_datetime(df["posted_date"])
    df["skill_list"] = df["skills"].apply(
        lambda value: [item.strip() for item in str(value).split(";") if item.strip()]
    )
    return df


def normalize_city_name(value: object) -> str:
    text = str(value).strip()
    if not text:
        return text
    return CITY_ALIAS_MAP.get(text.lower(), text)


def _standardize_jobs_df(df: pd.DataFrame) -> pd.DataFrame:
    required_cols = {"job_id", "role", "city", "posted_date", "title", "company", "skills"}
    missing = required_cols - set(df.columns)
    if missing:
        raise ValueError(f"Job dataset missing columns: {sorted(missing)}")
    df = df.copy()
    df["posted_date"] = pd.to_datetime(df["posted_date"], errors="coerce")
    df["city"] = df["city"].apply(normalize_city_name)
    df["skills"] = df["skills"].fillna("")
    df["skill_list"] = df["skills"].apply(
        lambda value: [item.strip() for item in str(value).split(";") if item.strip()]
    )
    return df.dropna(subset=["posted_date"])


@st.cache_data(show_spinner=False)
def load_cached_jobs() -> pd.DataFrame:
    base = load_jobs()
    if LIVE_CACHE_PATH.exists():
        try:
            live = pd.read_csv(LIVE_CACHE_PATH)
            combined = pd.concat([base, live], ignore_index=True)
            return _standardize_jobs_df(combined)
        except Exception:
            return _standardize_jobs_df(base)
    return _standardize_jobs_df(base)


def _load_snapshot_history() -> list[dict[str, object]]:
    if not SNAPSHOT_HISTORY_PATH.exists():
        return []
    try:
        return json.loads(SNAPSHOT_HISTORY_PATH.read_text(encoding="utf-8"))
    except Exception:
        return []


def _save_snapshot_history(entries: list[dict[str, object]]) -> None:
    SNAPSHOT_HISTORY_PATH.write_text(json.dumps(entries, indent=2), encoding="utf-8")


def build_profile_key(profile_context: str, role: str, city: str, github_username: str) -> str:
    fingerprint = f"{normalize_token(profile_context)}|{role}|{city}|{github_username}"
    return hashlib.sha256(fingerprint.encode("utf-8")).hexdigest()[:16]


def record_snapshot(profile_key: str, role: str, city: str, analysis: dict, compatibility_data: dict[str, int], student_skills: list[str]) -> None:
    history = _load_snapshot_history()
    today = pd.Timestamp.utcnow().strftime("%Y-%m-%d")
    entry = {
        "profile_key": profile_key,
        "date": today,
        "role": role,
        "city": city,
        "decay_risk": analysis["score"],
        "compatibility": compatibility_data["overall"],
        "missing_skills": analysis["missing"][:5],
        "detected_skills": student_skills[:12],
    }
    history = [item for item in history if not (item.get("profile_key") == profile_key and item.get("date") == today)]
    history.append(entry)
    history = sorted(history, key=lambda item: (item.get("profile_key", ""), item.get("date", "")))[-200:]
    _save_snapshot_history(history)


def get_profile_history(profile_key: str) -> pd.DataFrame:
    history = [item for item in _load_snapshot_history() if item.get("profile_key") == profile_key]
    if not history:
        return pd.DataFrame()
    df = pd.DataFrame(history).sort_values("date")
    return df


def _dummy_live_fetch(role: str, city: str, limit: int) -> pd.DataFrame:
    raise RuntimeError(
        "Live refresh is not configured. Provide a job API, or upload a live CSV, or use curated mode."
    )


def fetch_live_jobs(role: str, city: str, limit: int) -> pd.DataFrame:
    try:
        return fetch_adzuna_jobs(role=role, city=city, limit=limit)
    except Exception as primary_error:
        raise RuntimeError(str(primary_error)) from primary_error


def _skill_in_text(skill_token: str, normalized_text: str) -> bool:
    if len(skill_token) <= 3:
        return bool(re.search(r"(?:^|\s)" + re.escape(skill_token) + r"(?:\s|$)", normalized_text))
    return skill_token in normalized_text


def extract_skills(text: str) -> list[str]:
    normalized = normalize_token(text)
    found: list[str] = []
    for skill in SKILL_CATALOG:
        if _skill_in_text(normalize_token(skill), normalized):
            found.append(skill)
    for alias, canonical in SKILL_ALIAS.items():
        if _skill_in_text(alias, normalized) and canonical not in found:
            found.append(canonical)
    return sorted(found)


def count_skills(rows: pd.DataFrame) -> Counter:
    counter: Counter = Counter()
    for skills in rows["skill_list"]:
        counter.update(skills)
    return counter


def analyze_market(df: pd.DataFrame, student_skills: list[str]) -> dict:
    latest_date = df["posted_date"].max()
    recent_cutoff = latest_date - pd.Timedelta(days=30)
    recent_rows = df[df["posted_date"] >= recent_cutoff]
    previous_rows = df[df["posted_date"] < recent_cutoff]
    recent_counts = count_skills(recent_rows)
    previous_counts = count_skills(previous_rows)
    total_recent = max(len(recent_rows), 1)
    total_previous = max(len(previous_rows), 1)
    required_skills = [skill for skill, count in recent_counts.most_common(10) if count >= 3]
    matched = sorted([skill for skill in student_skills if skill in required_skills])
    missing = sorted([skill for skill in required_skills if skill not in student_skills])

    trend_rows = []
    for skill in sorted(set(recent_counts) | set(previous_counts)):
        recent_rate = recent_counts[skill] / total_recent
        previous_rate = previous_counts[skill] / total_previous
        delta = recent_rate - previous_rate
        trend_rows.append(
            {
                "skill": skill,
                "recent_mentions": recent_counts[skill],
                "previous_mentions": previous_counts[skill],
                "delta": round(delta, 3),
                "recent_rate": recent_rate,
            }
        )

    trend_df = pd.DataFrame(trend_rows).sort_values(by=["delta", "recent_mentions"], ascending=[False, False])
    rising = trend_df[trend_df["delta"] > 0].head(5)
    declining = trend_df.sort_values(by=["delta", "previous_mentions"]).head(5)
    gap_ratio = len(missing) / max(len(required_skills), 1)
    rising_missing = sum(1 for skill in rising["skill"].tolist() if skill in missing)
    trend_bonus = min(rising_missing * 8, 24)
    score = min(100, round(gap_ratio * 70 + trend_bonus + 10))

    explanations = []
    for _, row in rising.head(3).iterrows():
        explanations.append(
            f"`{row['skill']}` appears in **{int(row['recent_mentions'])}/{total_recent}** recent postings "
            f"vs **{int(row['previous_mentions'])}/{total_previous}** older postings (delta={row['delta']})."
        )

    return {
        "recent_rows": recent_rows,
        "previous_rows": previous_rows,
        "required_skills": required_skills,
        "matched": matched,
        "missing": missing,
        "rising": rising,
        "declining": declining,
        "score": score,
        "recent_cutoff": recent_cutoff,
        "explanations": explanations,
        "trend_df": trend_df,
    }


def compute_resume_compatibility(student_skills: list[str], required_skills: list[str], profile_text: str, role: str) -> dict:
    if required_skills:
        skill_match_pct = len([s for s in student_skills if s in required_skills]) / len(required_skills)
    else:
        skill_match_pct = 0.0
    all_market_skills = set(required_skills)
    keyword_hits = sum(1 for s in student_skills if s in all_market_skills)
    keyword_score = min(keyword_hits / max(len(all_market_skills), 1), 1.0)
    text_len = len(profile_text.strip())
    has_sections = any(kw in profile_text.lower() for kw in ["skills", "projects", "experience", "education", "profile"])
    has_numbers = bool(re.search(r"\d+", profile_text))
    completeness = min(1.0, (text_len / 800) * 0.5 + (0.3 if has_sections else 0) + (0.2 if has_numbers else 0))
    role_tokens = set(normalize_token(role).split())
    profile_norm = normalize_token(profile_text)
    role_hits = sum(1 for token in role_tokens if token in profile_norm)
    role_alignment = min(1.0, role_hits / max(len(role_tokens), 1))
    overall = round(skill_match_pct * 50 + keyword_score * 20 + completeness * 15 + role_alignment * 15)
    return {
        "overall": min(overall, 100),
        "skill_match": round(skill_match_pct * 100),
        "keyword_density": round(keyword_score * 100),
        "completeness": round(completeness * 100),
        "role_alignment": round(role_alignment * 100),
    }


def roadmap_for_skills(missing: list[str]) -> list[dict[str, str]]:
    default_tasks = [
        "Watch one focused tutorial and take concise notes.",
        "Implement one small hands-on exercise using sample data.",
        "Convert the exercise into a mini portfolio artifact.",
    ]
    roadmap = []
    selected = missing[:3] if missing else ["Power BI", "Dashboard Storytelling", "Prompt Engineering"]
    for index, skill in enumerate(selected, start=1):
        task = PROJECT_MAP.get(skill, default_tasks[(index - 1) % len(default_tasks)])
        roadmap.append({"day": f"Day {index * 2 - 1}-{index * 2}", "focus": skill, "task": task})
    roadmap.append({"day": "Day 7", "focus": "Proof Pack", "task": "Publish the mini project, write a README, and add one resume bullet with measurable impact."})
    return roadmap


def build_proof_pack(missing: list[str], student_skills: list[str]) -> dict[str, str]:
    target_skill = missing[0] if missing else "Power BI"
    support_skill = missing[1] if len(missing) > 1 else (student_skills[0] if student_skills else "SQL")
    project_title = f"{target_skill} Career Proof Dashboard"
    project_idea = PROJECT_MAP.get(target_skill, f"Build a mini project showing {target_skill} in action with a realistic student-friendly dataset.")
    resume_bullet = (
        f"Built a {target_skill}-focused analytics project integrating {support_skill} to solve a real reporting use case, "
        "documented insights, and packaged the work for recruiter review."
    )
    return {
        "title": project_title,
        "idea": project_idea,
        "resume_bullet": resume_bullet,
        "github_blurb": f"{project_title}: a portfolio project built to demonstrate market-relevant analytics skills for fresher roles.",
    }


def build_market_alert(analysis: dict, student_skills: list[str], target_role: str, portfolio_source: str) -> dict[str, str]:
    rising_missing = [skill for skill in analysis["rising"]["skill"].tolist() if skill not in student_skills]
    focus_skill = rising_missing[0] if rising_missing else (analysis["missing"][0] if analysis["missing"] else "")
    if not focus_skill:
        return {
            "title": "No critical alert right now",
            "body": f"Your current profile already covers the strongest visible signals for `{target_role}` in this dataset window.",
        }
    row = analysis["trend_df"][analysis["trend_df"]["skill"] == focus_skill].head(1)
    if row.empty:
        evidence = f"`{focus_skill}` is appearing more frequently in the recent market window."
    else:
        row = row.iloc[0]
        evidence = (
            f"`{focus_skill}` appears in {int(row['recent_mentions'])}/{len(analysis['recent_rows']) or 1} recent postings "
            f"vs {int(row['previous_mentions'])}/{len(analysis['previous_rows']) or 1} older postings."
        )
    source_line = f"Based on the provided profile context and portfolio input: `{portfolio_source}`." if portfolio_source else "Based on the provided profile context."
    return {
        "title": f"Market alert: `{focus_skill}` is becoming more important",
        "body": f"{evidence} You do not currently show this skill strongly in your profile. {source_line}",
    }


def generate_micro_curriculum(missing: list[str]) -> list[dict[str, object]]:
    selected = missing[:3] if missing else ["Power BI", "SQL", "Prompt Engineering"]
    curriculum = []
    for skill in selected:
        resources = LEARNING_RESOURCES.get(skill, [])
        lessons = MICRO_CURRICULUM_TEMPLATES.get(
            skill,
            [
                f"Study the core concepts behind {skill}.",
                f"Build one small hands-on exercise using {skill}.",
                f"Package the result into a visible proof-of-skill artifact.",
            ],
        )
        curriculum.append(
            {
                "skill": skill,
                "lessons": lessons,
                "resources": resources[:2],
            }
        )
    return curriculum


def generate_gemini_curriculum(missing: list[str], role: str, city: str, student_skills: list[str]) -> tuple[str | None, str | None]:
    api_key = get_secret("GEMINI_API_KEY")
    if not api_key:
        return None, "Gemini key not detected."
    model_name = get_secret("GEMINI_MODEL", "gemini-2.5-flash")
    focus_skills = missing[:3] if missing else ["Power BI", "SQL", "Prompt Engineering"]
    prompt = f"""
You are helping generate a concise, practical micro-curriculum for an early-career candidate.

Target role: {role}
City: {city}
Current detected skills: {", ".join(student_skills) if student_skills else "None"}
Missing high-demand skills: {", ".join(focus_skills)}

Create a 4-hour proof-oriented learning plan in markdown with these sections:
1. Goal
2. 3 learning steps
3. 1 mini project idea
4. 2 concrete portfolio outcomes

Keep it realistic, concise, and suitable for a student project demo.
""".strip()
    response = requests.post(
        f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent",
        params={"key": api_key},
        headers={"Content-Type": "application/json"},
        json={
            "contents": [
                {
                    "parts": [
                        {"text": prompt}
                    ]
                }
            ]
        },
        timeout=30,
    )
    if not response.ok:
        return None, f"Gemini API error {response.status_code} on `{model_name}`: {response.text[:220]}"
    data = response.json()
    candidates = data.get("candidates") or []
    if not candidates:
        return None, "Gemini returned no candidates."
    parts = candidates[0].get("content", {}).get("parts", [])
    text_parts = [part.get("text", "") for part in parts if part.get("text")]
    generated = "\n".join(text_parts).strip()
    if not generated:
        return None, "Gemini returned an empty response."
    return generated, None


def score_label(score: int) -> str:
    if score >= 75:
        return "High risk"
    if score >= 45:
        return "Medium risk"
    return "Low risk"


def build_radar_chart(student_skills: list[str], required_skills: list[str]) -> go.Figure:
    categories, student_vals, market_vals = [], [], []
    for category, category_skills in SKILL_CATEGORIES.items():
        student_count = sum(1 for skill in category_skills if skill in student_skills)
        market_count = sum(1 for skill in category_skills if skill in required_skills)
        total = max(len(category_skills), 1)
        categories.append(category)
        student_vals.append(round(student_count / total * 100))
        market_vals.append(round(market_count / total * 100))
    fig = go.Figure()
    fig.add_trace(go.Scatterpolar(r=student_vals + [student_vals[0]], theta=categories + [categories[0]], fill="toself", name="Your Skills", fillcolor="rgba(102, 126, 234, 0.15)", line=dict(color="#667eea", width=2)))
    fig.add_trace(go.Scatterpolar(r=market_vals + [market_vals[0]], theta=categories + [categories[0]], fill="toself", name="Market Demand", fillcolor="rgba(234, 102, 102, 0.10)", line=dict(color="#ea6666", width=2, dash="dash")))
    fig.update_layout(polar=dict(radialaxis=dict(visible=True, range=[0, 100])), showlegend=True, height=400, margin=dict(l=60, r=60, t=40, b=40))
    return fig


def build_trend_chart(analysis: dict) -> go.Figure:
    trend_df = analysis["trend_df"].head(12).copy()
    if trend_df.empty:
        return go.Figure()
    trend_df = trend_df.sort_values("delta", ascending=True)
    fig = go.Figure()
    fig.add_trace(go.Bar(y=trend_df["skill"], x=trend_df["previous_mentions"], name="Previous Period", orientation="h", marker_color="#94a3b8"))
    fig.add_trace(go.Bar(y=trend_df["skill"], x=trend_df["recent_mentions"], name="Recent Period", orientation="h", marker_color="#667eea"))
    fig.update_layout(barmode="group", height=max(300, len(trend_df) * 35), margin=dict(l=10, r=10, t=30, b=10), xaxis_title="Mentions in Job Postings")
    return fig


def build_skill_gap_chart(matched: list[str], missing: list[str]) -> go.Figure:
    skills = matched + missing
    if not skills:
        return go.Figure()
    colors = ["#28a745"] * len(matched) + ["#dc3545"] * len(missing)
    labels = ["Matched"] * len(matched) + ["Missing"] * len(missing)
    fig = go.Figure(go.Bar(y=skills, x=[1] * len(skills), orientation="h", marker_color=colors, text=labels, textposition="inside"))
    fig.update_layout(height=max(250, len(skills) * 32), margin=dict(l=10, r=10, t=10, b=10), xaxis=dict(visible=False), showlegend=False)
    return fig


def build_resume_gauge(score: int) -> go.Figure:
    color = "#28a745" if score >= 70 else "#ffc107" if score >= 40 else "#dc3545"
    fig = go.Figure(go.Indicator(
        mode="gauge+number",
        value=score,
        title={"text": "Resume Compatibility", "font": {"size": 16}},
        number={"suffix": "%", "font": {"size": 34}},
        gauge={
            "axis": {"range": [0, 100], "tickwidth": 1},
            "bar": {"color": color, "thickness": 0.3},
            "steps": [
                {"range": [0, 40], "color": "#f8d7da"},
                {"range": [40, 70], "color": "#fff3cd"},
                {"range": [70, 100], "color": "#d4edda"},
            ],
        },
    ))
    fig.update_layout(height=250, margin=dict(l=20, r=20, t=40, b=20))
    return fig


def build_career_path_chart(role: str, student_skills: list[str]) -> go.Figure:
    path_data = next((CAREER_PATHS[key] for key in CAREER_PATHS if key.lower() in role.lower() or role.lower() in key.lower()), list(CAREER_PATHS.values())[0])
    labels, parents, values, colors = [path_data["current"]], [""], [100], ["#667eea"]
    for path in path_data["paths"]:
        labels.append(f"{path['role']} ({path['years']})")
        parents.append(path_data["current"])
        skill_match = sum(1 for skill in path["key_skills"] if skill in student_skills)
        readiness = round(skill_match / max(len(path["key_skills"]), 1) * 100)
        values.append(readiness)
        colors.append("#28a745" if readiness >= 70 else "#ffc107" if readiness >= 40 else "#dc3545")
    fig = go.Figure(go.Treemap(labels=labels, parents=parents, values=[max(value, 10) for value in values], marker=dict(colors=colors), textinfo="label+text", text=[f"Readiness: {value}%" if index > 0 else "You are here" for index, value in enumerate(values)]))
    fig.update_layout(height=350, margin=dict(l=10, r=10, t=30, b=10))
    return fig


def build_history_chart(history_df: pd.DataFrame) -> go.Figure:
    fig = go.Figure()
    if history_df.empty:
        return fig
    fig.add_trace(go.Scatter(x=history_df["date"], y=history_df["compatibility"], mode="lines+markers", name="Resume Compatibility", line=dict(color="#667eea", width=3)))
    fig.add_trace(go.Scatter(x=history_df["date"], y=history_df["decay_risk"], mode="lines+markers", name="Skill Decay Risk", line=dict(color="#dc3545", width=3)))
    fig.update_layout(height=300, margin=dict(l=10, r=10, t=30, b=10), yaxis_title="Score", xaxis_title="Snapshot Date")
    return fig


def _pdf_safe(text: str) -> str:
    text = re.sub(r"[`*]", "", text)
    text = text.replace("\u0394", "delta").replace("\u2014", "-").replace("\u2013", "-")
    text = text.replace("\u2018", "'").replace("\u2019", "'").replace("\u201c", '"').replace("\u201d", '"')
    return text.encode("latin-1", errors="replace").decode("latin-1")


def generate_pdf_report(role: str, city: str, filtered_jobs: pd.DataFrame, analysis: dict, roadmap: list[dict[str, str]], proof_pack: dict[str, str], compatibility_data: dict[str, int], student_skills: list[str]) -> bytes | None:
    if FPDF is None:
        return None
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    def write_block(text: str, size: int = 10) -> None:
        pdf.set_font("Helvetica", "", size)
        pdf.multi_cell(w=0, h=6, text=_pdf_safe(text), new_x="LMARGIN", new_y="NEXT")

    def heading(text: str) -> None:
        pdf.set_font("Helvetica", "B", 13)
        pdf.cell(w=0, h=10, text=text, new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("Helvetica", "B", 22)
    pdf.cell(w=0, h=15, text="SkillPulse Report", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(w=0, h=8, text=_pdf_safe(f"Role: {role}  |  City: {city}  |  Postings Analyzed: {len(filtered_jobs)}"), align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)
    heading("Scores")
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(w=0, h=7, text=f"Skill Decay Risk: {analysis['score']}/100 ({score_label(analysis['score'])})", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(w=0, h=7, text=f"Resume Compatibility: {compatibility_data['overall']}%", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    heading("Matched Skills")
    write_block(", ".join(analysis["matched"]) if analysis["matched"] else "None")
    heading("Missing High-Demand Skills")
    write_block(", ".join(analysis["missing"]) if analysis["missing"] else "None")
    heading("Market Trend Insights")
    for explanation in analysis["explanations"]:
        write_block(f"- {explanation}")
    heading("7-Day Micro Roadmap")
    for item in roadmap:
        write_block(f"{item['day']} - {item['focus']}: {item['task']}")
    heading("Proof Pack")
    write_block(f"Project: {proof_pack['title']}")
    write_block(f"Idea: {proof_pack['idea']}")
    write_block(f"Resume Bullet: {proof_pack['resume_bullet']}")
    heading("Your Skills")
    write_block(", ".join(student_skills) if student_skills else "None detected")
    heading("Resume Compatibility Breakdown")
    for label, key in [("Skill Match", "skill_match"), ("Keyword Density", "keyword_density"), ("Completeness", "completeness"), ("Role Alignment", "role_alignment")]:
        pdf.cell(w=0, h=6, text=f"{label}: {compatibility_data[key]}%", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def display_skill_tags_html(skills: list[str], css_class: str, empty_text: str) -> None:
    if skills:
        tags = "".join(f'<span class="skill-tag {css_class}">{skill}</span>' for skill in skills)
        st.markdown(tags, unsafe_allow_html=True)
    else:
        st.caption(empty_text)


def main() -> None:
    st.set_page_config(page_title="SkillPulse", page_icon="SP", layout="wide")
    st.markdown(CUSTOM_CSS, unsafe_allow_html=True)
    st.title("SkillPulse")
    st.caption("Detect skill decay, compare against current job signals, and generate a proof-oriented career roadmap.")

    jobs_df = load_cached_jobs()

    with st.sidebar:
        st.header("Configuration")
        role = st.selectbox("Target role", sorted(jobs_df["role"].unique()))
        city = st.selectbox("Target city", sorted(jobs_df["city"].unique()))
        st.divider()
        st.subheader("Resume Upload")
        resume_upload = st.file_uploader("Upload your resume (PDF, DOCX, or TXT)", type=["pdf", "docx", "txt"], help="We'll extract text and auto-detect your skills.")
        uploaded_resume_text = ""
        if resume_upload is not None:
            with st.spinner("Extracting text from resume..."):
                uploaded_resume_text = extract_text_from_upload(resume_upload)
            if uploaded_resume_text and not uploaded_resume_text.startswith("["):
                st.success(f"Extracted {len(uploaded_resume_text)} characters from {resume_upload.name}")
            elif uploaded_resume_text.startswith("["):
                st.warning(uploaded_resume_text)
                uploaded_resume_text = ""
        st.divider()
        st.subheader("Student Profile")
        profile_choice = st.radio(
            "Profile source",
            options=[
                "Uploaded Resume" if resume_upload else "Upload a resume above",
                "Fresh profile",
                "Sample (Strong match)",
                "Sample (Weak match)",
                "Resume (Visshva - AIML/Data)",
                "Resume (Visshva - SDE/Full-stack)",
                "Custom paste",
            ],
            index=1,
        )
        if "Uploaded Resume" in profile_choice and uploaded_resume_text:
            profile_text = uploaded_resume_text
        elif profile_choice == "Fresh profile":
            profile_text = ""
        elif profile_choice == "Sample (Strong match)":
            profile_text = SAMPLE_RESUME_PATH.read_text(encoding="utf-8")
        elif profile_choice == "Sample (Weak match)":
            profile_text = SAMPLE_RESUME_ALT_PATH.read_text(encoding="utf-8")
        elif profile_choice == "Resume (Visshva - AIML/Data)":
            profile_text = RESUME_VISSHVA_AIML_PATH.read_text(encoding="utf-8")
        elif profile_choice == "Resume (Visshva - SDE/Full-stack)":
            profile_text = RESUME_VISSHVA_SDE_PATH.read_text(encoding="utf-8")
        else:
            profile_text = ""
        profile_text = st.text_area(
            "Edit or paste profile text",
            value=profile_text,
            height=220,
            placeholder="Paste a fresh resume summary, skills, projects, or achievements here...",
        )
        portfolio_source = st.text_input(
            "GitHub / portfolio URL or notes (optional)",
            placeholder="https://github.com/username or brief portfolio notes",
        )
        github_username = parse_github_username(portfolio_source)
        github_data = {"username": "", "skills": [], "summary": "", "repos": []}
        if github_username:
            try:
                with st.spinner("Analyzing GitHub profile..."):
                    github_data = fetch_github_profile_data(github_username)
                st.success(f"Loaded public GitHub signals for @{github_username}")
                if github_data["skills"]:
                    st.caption("GitHub-detected skills: " + ", ".join(github_data["skills"][:8]))
            except Exception as error:
                st.warning(f"GitHub analysis unavailable: {error}")
        st.divider()
        st.subheader("Data Source")
        uploaded_csv = st.file_uploader("Upload a live jobs CSV (optional)", type=["csv"])
        if uploaded_csv is not None:
            try:
                jobs_df = _standardize_jobs_df(pd.read_csv(uploaded_csv))
                st.success("Loaded uploaded job dataset.")
            except Exception as error:
                st.error(f"Could not read uploaded CSV: {error}")
        live_limit = st.slider("Live refresh size", min_value=10, max_value=80, value=30, step=10)
        if st.button("Refresh live signals"):
            try:
                live_df = _standardize_jobs_df(fetch_live_jobs(role=role, city=city, limit=live_limit))
                live_df.to_csv(LIVE_CACHE_PATH, index=False)
                st.success("Live refresh complete.")
                st.cache_data.clear()
            except Exception as error:
                st.warning(f"Live refresh not available: {error}")

    filtered_jobs = jobs_df[(jobs_df["role"] == role) & (jobs_df["city"] == city)].copy()
    if filtered_jobs.empty:
        st.error("No job postings available for this role/city combination in the current dataset.")
        st.stop()

    github_profile_text = ""
    if github_data["summary"]:
        github_profile_text += f"\nGitHub bio: {github_data['summary']}"
    if github_data["repos"]:
        github_profile_text += "\nRecent repositories:\n- " + "\n- ".join(github_data["repos"])
    if github_data["skills"]:
        github_profile_text += "\nGitHub-detected skills: " + ", ".join(github_data["skills"])

    profile_context = profile_text
    if portfolio_source.strip():
        profile_context = f"{profile_text}\n{portfolio_source.strip()}"
    if github_profile_text.strip():
        profile_context = f"{profile_context}\n{github_profile_text}".strip()

    student_skills = extract_skills(profile_context)
    analysis = analyze_market(filtered_jobs, student_skills)
    roadmap = roadmap_for_skills(analysis["missing"])
    proof_pack = build_proof_pack(analysis["missing"], student_skills)
    compatibility_data = compute_resume_compatibility(student_skills, analysis["required_skills"], profile_context, role)
    market_alert = build_market_alert(analysis, student_skills, role, portfolio_source.strip())
    micro_curriculum = generate_micro_curriculum(analysis["missing"])
    gemini_curriculum = None
    gemini_status = "Gemini not checked yet."
    try:
        gemini_curriculum, gemini_error = generate_gemini_curriculum(analysis["missing"], role, city, student_skills)
        gemini_status = "Gemini live generation enabled." if gemini_curriculum else (gemini_error or "Gemini did not return content.")
    except Exception as error:
        gemini_curriculum = None
        gemini_status = f"Gemini request failed: {error}"
    profile_key = build_profile_key(profile_context, role, city, github_data.get("username", ""))
    record_snapshot(profile_key, role, city, analysis, compatibility_data, student_skills)
    history_df = get_profile_history(profile_key)

    metrics = st.columns(4)
    metrics[0].metric("Job Postings Analyzed", len(filtered_jobs))
    metrics[1].metric("Skills Detected", len(student_skills))
    metrics[2].metric("Market Matches", len(analysis["matched"]))
    metrics[3].metric("Skill Decay Risk", f"{analysis['score']}/100", score_label(analysis["score"]))
    st.caption(f"Trend window: recent postings since **{analysis['recent_cutoff'].date()}** vs older postings.")

    tab1, tab2, tab3, tab4, tab5, tab6, tab7 = st.tabs([
        "Dashboard",
        "Market Trends",
        "Resume Compatibility",
        "Roadmap & Proof Pack",
        "Micro-Curriculum",
        "Learning Resources",
        "Career Paths",
    ])

    with tab1:
        left, right = st.columns([1.1, 0.9])
        with left:
            st.markdown('<p class="section-header">Skill Radar - You vs Market</p>', unsafe_allow_html=True)
            st.plotly_chart(build_radar_chart(student_skills, analysis["required_skills"]), use_container_width=True)
            st.markdown('<p class="section-header">Skill Gap Analysis</p>', unsafe_allow_html=True)
            gap_fig = build_skill_gap_chart(analysis["matched"], analysis["missing"])
            if analysis["matched"] or analysis["missing"]:
                st.plotly_chart(gap_fig, use_container_width=True)
        with right:
            st.markdown('<p class="section-header">Your Skills</p>', unsafe_allow_html=True)
            display_skill_tags_html(student_skills, "skill-matched", "No recognizable skills found.")
            st.markdown('<p class="section-header">Matched with Market</p>', unsafe_allow_html=True)
            display_skill_tags_html(analysis["matched"], "skill-matched", "No strong market matches yet.")
            st.markdown('<p class="section-header">Missing High-Demand Skills</p>', unsafe_allow_html=True)
            display_skill_tags_html(analysis["missing"], "skill-missing", "No major market gaps detected.")
            st.markdown('<p class="section-header">Why the Market is Shifting</p>', unsafe_allow_html=True)
            for line in analysis["explanations"]:
                st.markdown(f"- {line}")

    with tab2:
        st.markdown('<p class="section-header">Skill Demand - Recent vs Previous Period</p>', unsafe_allow_html=True)
        st.plotly_chart(build_trend_chart(analysis), use_container_width=True)
        col_rise, col_decay = st.columns(2)
        with col_rise:
            st.markdown("**Rising Skills**")
            rising_df = analysis["rising"][["skill", "recent_mentions", "previous_mentions", "delta"]].copy()
            rising_df.columns = ["Skill", "Recent", "Previous", "Trend Delta"]
            st.dataframe(rising_df, use_container_width=True, hide_index=True)
        with col_decay:
            st.markdown("**Declining Skills**")
            decline_df = analysis["declining"][["skill", "recent_mentions", "previous_mentions", "delta"]].copy()
            decline_df.columns = ["Skill", "Recent", "Previous", "Trend Delta"]
            st.dataframe(decline_df, use_container_width=True, hide_index=True)

    with tab3:
        st.markdown('<p class="section-header">Resume Compatibility Analysis</p>', unsafe_allow_html=True)
        compat_left, compat_right = st.columns([1, 1.2])
        with compat_left:
            st.plotly_chart(build_resume_gauge(compatibility_data["overall"]), use_container_width=True)
            if compatibility_data["overall"] >= 70:
                st.success("Strong resume compatibility. Your profile aligns well with market requirements.")
            elif compatibility_data["overall"] >= 40:
                st.warning("Moderate resume compatibility. Adding missing skills would improve your alignment.")
            else:
                st.error("Low resume compatibility. Focus on acquiring the missing high-demand skills listed.")
        with compat_right:
            st.markdown("**Score Breakdown**")
            for label, key, caption in [
                ("Skill Match", "skill_match", "How many market-required skills you have."),
                ("Keyword Density", "keyword_density", "Coverage of in-demand keywords in your profile."),
                ("Profile Completeness", "completeness", "Structure, detail, and quantifiable achievements."),
                ("Role Alignment", "role_alignment", "How well your profile targets this specific role."),
            ]:
                st.markdown(f"**{label}** — `{compatibility_data[key]}%`")
                st.progress(compatibility_data[key] / 100)
                st.caption(caption)
            if github_data["username"]:
                st.markdown("**GitHub Signal Overlay**")
                st.caption("Public GitHub repositories and topics are folded into the profile context to make the analysis closer to a portfolio-aware evaluation.")
                st.write(f"GitHub profile analyzed: `@{github_data['username']}`")

    with tab4:
        road_left, road_right = st.columns(2)
        with road_left:
            st.markdown('<p class="section-header">7-Day Micro Roadmap</p>', unsafe_allow_html=True)
            for item in roadmap:
                st.markdown(f"**{item['day']} - {item['focus']}**")
                st.write(item["task"])
                st.divider()
        with road_right:
            st.markdown('<p class="section-header">Proof Pack</p>', unsafe_allow_html=True)
            st.markdown(f"**Project Title:** `{proof_pack['title']}`")
            st.info(proof_pack["idea"])
            st.success(proof_pack["resume_bullet"])
            st.code(proof_pack["github_blurb"], language=None)
        st.markdown('<p class="section-header">Download Report</p>', unsafe_allow_html=True)
        report_md = "\n".join([
            "# SkillPulse Report",
            f"- Role: {role}",
            f"- City: {city}",
            f"- Job postings analyzed: {len(filtered_jobs)}",
            f"- Skill decay risk: {analysis['score']}/100 ({score_label(analysis['score'])})",
            f"- Resume compatibility: {compatibility_data['overall']}%",
            "",
            "## Your Skills",
            ", ".join(student_skills) if student_skills else "None detected",
            "",
            "## Matched Skills",
            ", ".join(analysis["matched"]) if analysis["matched"] else "None",
            "",
            "## Missing High-Demand Skills",
            ", ".join(analysis["missing"]) if analysis["missing"] else "None",
            "",
            "## Top Trending Signals",
            "\n".join([f"- {line}" for line in analysis["explanations"]]) if analysis["explanations"] else "None",
            "",
            "## Resume Compatibility Breakdown",
            f"- Overall: {compatibility_data['overall']}%",
            f"- Skill Match: {compatibility_data['skill_match']}%",
            f"- Keyword Density: {compatibility_data['keyword_density']}%",
            f"- Completeness: {compatibility_data['completeness']}%",
            f"- Role Alignment: {compatibility_data['role_alignment']}%",
            "",
            "## 7-Day Roadmap",
            "\n".join([f"- **{item['day']}**: {item['focus']} — {item['task']}" for item in roadmap]),
            "",
            "## Proof Pack",
            f"- Title: {proof_pack['title']}",
            f"- Idea: {proof_pack['idea']}",
            f"- Resume Bullet: {proof_pack['resume_bullet']}",
            f"- GitHub Blurb: {proof_pack['github_blurb']}",
        ])
        dl_col1, dl_col2 = st.columns(2)
        with dl_col1:
            st.download_button("Download Report (Markdown)", data=report_md.encode("utf-8"), file_name="skillpulse_report.md", mime="text/markdown")
        with dl_col2:
            pdf_bytes = generate_pdf_report(role, city, filtered_jobs, analysis, roadmap, proof_pack, compatibility_data, student_skills)
            if pdf_bytes:
                st.download_button("Download Report (PDF)", data=pdf_bytes, file_name="skillpulse_report.pdf", mime="application/pdf")
            else:
                st.caption("Install `fpdf2` for PDF report downloads.")

    with tab5:
        st.markdown('<p class="section-header">Continuous Skill Alert</p>', unsafe_allow_html=True)
        st.warning(market_alert["title"])
        st.write(market_alert["body"])
        st.markdown('<p class="section-header">Profile Snapshot History</p>', unsafe_allow_html=True)
        if len(history_df) >= 1:
            st.plotly_chart(build_history_chart(history_df), use_container_width=True)
            latest = history_df.iloc[-1]
            if len(history_df) > 1:
                previous = history_df.iloc[-2]
                st.caption(
                    f"Latest snapshot: compatibility {latest['compatibility']}% ({latest['compatibility'] - previous['compatibility']:+} vs previous), "
                    f"decay risk {latest['decay_risk']} ({latest['decay_risk'] - previous['decay_risk']:+} vs previous)."
                )
            else:
                st.caption("This is the first saved snapshot for this profile-role combination. Revisit after edits or skill updates to show progress over time.")
        st.markdown('<p class="section-header">Micro-Curriculum Generator</p>', unsafe_allow_html=True)
        st.caption("A short, proof-oriented curriculum generated from the strongest missing market signals.")
        for block in micro_curriculum:
            with st.expander(block["skill"], expanded=block["skill"] == micro_curriculum[0]["skill"]):
                for idx, lesson in enumerate(block["lessons"], start=1):
                    st.markdown(f"{idx}. {lesson}")
                if block["resources"]:
                    st.markdown("**Suggested resources**")
                    for resource in block["resources"]:
                        st.markdown(f"- [{resource['title']}]({resource['url']}) — {resource['type']} ({resource['time']})")
        st.markdown('<p class="section-header">Gemini Vision Layer</p>', unsafe_allow_html=True)
        st.caption(gemini_status)
        if gemini_curriculum:
            st.markdown(gemini_curriculum)
        else:
            st.caption("Gemini-enhanced curriculum is optional. Add `GEMINI_API_KEY` in Streamlit secrets or environment variables to enable live generation.")

    with tab6:
        st.markdown('<p class="section-header">Recommended Learning Resources</p>', unsafe_allow_html=True)
        target_skills = analysis["missing"] if analysis["missing"] else list(LEARNING_RESOURCES.keys())[:3]
        for skill in target_skills:
            resources = LEARNING_RESOURCES.get(skill)
            if resources:
                with st.expander(skill, expanded=skill == target_skills[0]):
                    for resource in resources:
                        st.markdown(f"**{resource['title']}** — _{resource['type']}_ ({resource['time']})  \n[Open Resource]({resource['url']})")
                        st.divider()
            else:
                with st.expander(skill):
                    st.caption(f"No curated resources yet for {skill}.")

    with tab7:
        st.markdown('<p class="section-header">Career Path Progression</p>', unsafe_allow_html=True)
        st.plotly_chart(build_career_path_chart(role, student_skills), use_container_width=True)

    with st.expander("Compare Two Profiles"):
        comp_col1, comp_col2 = st.columns(2)
        with comp_col1:
            profile_a = st.text_area("Profile A", height=180, key="comp_a")
        with comp_col2:
            profile_b = st.text_area("Profile B", height=180, key="comp_b")
        if profile_a.strip() and profile_b.strip():
            skills_a = extract_skills(profile_a)
            skills_b = extract_skills(profile_b)
            analysis_a = analyze_market(filtered_jobs, skills_a)
            analysis_b = analyze_market(filtered_jobs, skills_b)
            res_col1, res_col2 = st.columns(2)
            with res_col1:
                st.markdown("**Profile A**")
                st.metric("Skills Detected", len(skills_a))
                st.metric("Market Matches", len(analysis_a["matched"]))
                st.metric("Decay Risk", f"{analysis_a['score']}/100")
            with res_col2:
                st.markdown("**Profile B**")
                st.metric("Skills Detected", len(skills_b))
                st.metric("Market Matches", len(analysis_b["matched"]))
                st.metric("Decay Risk", f"{analysis_b['score']}/100")

    with st.expander("Real-World Relevance"):
        st.markdown("""
        - **Students:** identify which current skills still align with hiring demand and which ones need reinforcement.
        - **Placement cells:** use one consistent market-signal view to guide multiple students.
        - **Early-career job seekers:** convert missing skills into focused portfolio evidence instead of generic learning plans.
        """)

    with st.expander("Reference Case Studies"):
        st.markdown("""
        - **Case 1:** a final-year student with `Excel`, `SQL`, and `Python` sees that `Power BI` and `Dashboard Storytelling` are becoming stronger differentiators.
        - **Case 2:** a weakly aligned student with only reporting experience gets a higher risk score and a targeted recovery roadmap.
        - **Case 3:** a placement mentor uses the same role dataset to guide multiple students consistently.
        """)

    with st.expander("SDG Alignment"):
        st.markdown("""
        - **SDG 4 (Quality Education):** guides students to targeted micro-skills and proof-based learning.
        - **SDG 8 (Decent Work & Economic Growth):** improves employability by aligning skills with real job demand.
        """)

    with st.expander("Evidence & Sources"):
        st.markdown("""
        - **Business analysis + agentic workflow trends (2026):** ModernAnalyst article on evolving BA responsibilities.
        - **Skills taxonomy / skills graph concept:** Cornerstone Skills Graph writeup.
        - **Workforce upskilling pressure:** Hays skills reporting and related workforce-skilling references.
        - **Important:** the app's analysis comes from the job dataset; these documents support the product rationale.
        """)


if __name__ == "__main__":
    main()
